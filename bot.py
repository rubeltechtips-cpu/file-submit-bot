import logging
import os
import io
import time
from telegram import (
    Update, ReplyKeyboardMarkup, KeyboardButton, InputFile,
    ReplyKeyboardRemove, InlineKeyboardMarkup, InlineKeyboardButton
)
from telegram.ext import (
    ApplicationBuilder, CommandHandler, MessageHandler,
    filters, ContextTypes, ConversationHandler, CallbackQueryHandler
)
from telegram.constants import ChatMemberStatus
from openpyxl import Workbook, load_workbook
import json

# ================ CONFIG ================
BOT_TOKEN = "8349208659:AAEyJikjx1tUri_PztFGRca_lPT0WilJ0N0"
ADMIN_ID = 8061006207
ADMIN_USERNAME = "Rubel_QSB"
CHANNEL_USERNAME = "quick_sell_bd"
DATA_DIR = "categories"

os.makedirs(DATA_DIR, exist_ok=True)

# ================ STATES ================
(
    BUY_MENU,
    BUY_SUB_MENU,
    ADMIN_PANEL,
    ADD_MAIN_CAT,
    REMOVE_MAIN_CAT,
    MANAGE_CATEGORY,
    MANAGE_SUB_CATEGORY,
    ADD_SUB_CAT,
    REMOVE_SUB_CAT,
    ADD_ITEMS,
    EDIT_PAYMENT,
    EDIT_PRICE_MAIN,
    EDIT_PRICE_SUB,
    RECEIVE_NEW_PRICE,
    GET_QUANTITY,
    WAIT_SCREENSHOT,
    DEPOSIT,
    GET_DEPOSIT_AMOUNT,
    DASHBOARD,
    SEND_NOTICE,
    VIEW_USER_PROFILE,
    SEARCH_USER_PROFILE,
    MANAGE_PAYMENT_CATEGORIES,
    # --- NEW STATES FOR BALANCE EDIT ---
    SEARCH_USER_FOR_BALANCE,
    BALANCE_EDIT_ACTION,
    RECEIVE_BALANCE_EDIT_AMOUNT
) = range(26)

# ================ LOGGER ================
logging.basicConfig(format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

# ================ DATA ================
# New: Hierarchical categories and prices
categories = {}
prices = {}
payment_info = "Bkash: 017XXXXXXXX\nNagad: 018XXXXXXXX\nBinance: yourmail@gmail.com"
# New: User balances dictionary
balances = {}
user_sales = {}
user_deposits = {}
user_info = {} # To store user ID and username mapping
# New: Dashboard metrics
total_deposits = 0
total_sales = 0
sales_count_per_category = {}
transaction_log = []
dashboard_message = "স্বাগতম! এটি আপনার বটের ড্যাশবোর্ড।"

# List of categories that require manual delivery
MANUAL_DELIVERY_CATEGORIES = []

# --- User data persistence ---
def load_user_data():
    global balances, user_sales, user_deposits, user_info, total_deposits, total_sales, transaction_log, categories, prices, sales_count_per_category, MANUAL_DELIVERY_CATEGORIES
    try:
        with open("user_data.json", "r", encoding='utf-8') as f:
            data = json.load(f)
            balances.update({int(k): v for k, v in data.get("balances", {}).items()})
            user_sales.update({int(k): v for k, v in data.get("user_sales", {}).items()})
            user_deposits.update({int(k): v for k, v in data.get("user_deposits", {}).items()})
            user_info.update({int(k): v for k, v in data.get("user_info", {}).items()})
            total_deposits = data.get("total_deposits", 0)
            total_sales = data.get("total_sales", 0)
            transaction_log = data.get("transaction_log", [])
            categories.update(data.get("categories", {}))
            prices.update(data.get("prices", {}))
            sales_count_per_category.update(data.get("sales_count_per_category", {}))
            MANUAL_DELIVERY_CATEGORIES = data.get("manual_delivery_categories", [])

    except FileNotFoundError:
        pass

def save_user_data():
    with open("user_data.json", "w", encoding='utf-8') as f:
        data = {
            "balances": balances,
            "user_sales": user_sales,
            "user_deposits": user_deposits,
            "user_info": user_info,
            "total_deposits": total_deposits,
            "total_sales": total_sales,
            "transaction_log": transaction_log,
            "categories": categories,
            "prices": prices,
            "sales_count_per_category": sales_count_per_category,
            "manual_delivery_categories": MANUAL_DELIVERY_CATEGORIES
        }
        json.dump(data, f, ensure_ascii=False, indent=4)
        
load_user_data()

# ================ HELPERS ================
def get_excel_path(main_cat: str, sub_cat: str) -> str:
    # Use a file name that combines both main and sub categories
    file_name = f"{main_cat}_{sub_cat}.xlsx".replace(" ", "_").replace("-", "_")
    return os.path.join(DATA_DIR, file_name)

def ensure_excel(main_cat: str, sub_cat: str):
    path = get_excel_path(main_cat, sub_cat)
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Items"
        wb.save(path)

def add_item_to_excel(main_cat: str, sub_cat: str, item: str):
    ensure_excel(main_cat, sub_cat)
    wb = load_workbook(get_excel_path(main_cat, sub_cat))
    ws = wb.active
    ws.append([item])
    wb.save(get_excel_path(main_cat, sub_cat))

def pop_items_from_excel(main_cat: str, sub_cat: str, qty: int):
    path = get_excel_path(main_cat, sub_cat)
    if not os.path.exists(path):
        return []
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    items = [r[0] for r in rows if r[0]]
    if len(items) < qty:
        return []
    
    result = items[:qty]
    new_items = items[qty:]
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    for item in new_items:
        ws.append([item])
    wb.save(path)
    
    return result

def count_items(main_cat: str, sub_cat: str) -> int:
    path = get_excel_path(main_cat, sub_cat)
    if not os.path.exists(path):
        return 0
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return len([r[0] for r in rows if r[0]])

def get_total_stock(main_cat: str) -> int:
    total = 0
    if main_cat in categories:
        for sub_cat in categories[main_cat]:
            total += count_items(main_cat, sub_cat)
    return total

def create_xlsx_file(items: list, file_name: str) -> io.BytesIO:
    """Creates an in-memory XLSX file from a list of items."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    for item in items:
        ws.append([item])
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def get_report_summary(transactions, days):
    """Calculates total deposits and sales for a given number of days."""
    end_timestamp = time.time()
    start_timestamp = end_timestamp - (days * 24 * 60 * 60)
    
    daily_deposits = 0
    daily_sales = 0
    
    for trans_type, _, amount, timestamp in transactions:
        if start_timestamp <= timestamp <= end_timestamp:
            if trans_type == 'deposit':
                daily_deposits += amount
            elif trans_type == 'sale':
                daily_sales += amount
    return daily_deposits, daily_sales
    
def get_user_transactions(user_id, transactions):
    """Filters transactions for a specific user."""
    return [t for t in transactions if t[1] == user_id]

# ================ HANDLERS ================
async def check_subscription(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    try:
        chat_member = await context.bot.get_chat_member(chat_id=f"@{CHANNEL_USERNAME}", user_id=user_id)
        if chat_member.status in [ChatMemberStatus.MEMBER, ChatMemberStatus.ADMINISTRATOR, ChatMemberStatus.OWNER]:
            return True
        else:
            keyboard = InlineKeyboardMarkup([
                [InlineKeyboardButton("Join Channel", url=f"https://t.me/{CHANNEL_USERNAME}")]
            ])
            await update.message.reply_text(
                "❌ আপনি এখনো আমাদের চ্যানেলে জয়েন করেননি।\n"
                "বট ব্যবহার করার জন্য অনুগ্রহ করে নিচের বাটনে ক্লিক করে চ্যানেলে জয়েন করুন।",
                reply_markup=keyboard
            )
            return False
    except Exception as e:
        logger.error(f"Error checking subscription: {e}")
        await update.message.reply_text("চ্যানেল সদস্যতা পরীক্ষা করতে সমস্যা হচ্ছে। অনুগ্রহ করে পরে আবার চেষ্টা করুন।")
        return False

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END
    
    user_id = update.effective_user.id
    username = update.effective_user.username
    first_name = update.effective_user.first_name
    
    # Store user info if it's new
    if user_id not in user_info:
        user_info[user_id] = {
            "username": username,
            "first_name": first_name,
            "last_name": update.effective_user.last_name,
            "id": user_id
        }
        save_user_data()

    current_balance = balances.get(user_id, 0)

    keyboard = [
        [KeyboardButton("🛒 Buy"), KeyboardButton("💰 Balance")],
        [KeyboardButton("💸 Deposit"), KeyboardButton("📞 Help")],
    ]
    if update.effective_user.id == ADMIN_ID:
        keyboard.append([KeyboardButton("⚙️ Admin Panel")])
    
    reply_markup = ReplyKeyboardMarkup(keyboard, resize_keyboard=True, one_time_keyboard=False)
    await update.message.reply_text(f"👋 Welcome! আপনার বর্তমান ব্যালেন্স: {current_balance} টাকা।", reply_markup=reply_markup)
    return ConversationHandler.END

async def menu_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END

    text = update.message.text
    
    if text == "🛒 Buy":
        if not categories:
            await update.message.reply_text("⚠️ এখন কোনো category নেই।")
            return ConversationHandler.END
        
        keyboard = []
        for cat in categories.keys():
            keyboard.append([KeyboardButton(cat)])
        
        keyboard.append([KeyboardButton("🔙 Back to Main Menu")])
        await update.message.reply_text("🛒 ক্যাটাগরি বেছে নিন:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return BUY_MENU

    if text == "💰 Balance":
        user_id = update.effective_user.id
        current_balance = balances.get(user_id, 0)
        await update.message.reply_text(f"আপনার বর্তমান ব্যালেন্স: {current_balance} টাকা।")
        return ConversationHandler.END

    if text == "💸 Deposit":
        await update.message.reply_text("আপনি কত টাকা ডিপোজিট করতে চান তা লিখুন:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
        return DEPOSIT

    if text == "📞 Help":
        keyboard = InlineKeyboardMarkup([
            [InlineKeyboardButton("📞 Chat with Admin", url=f"tg://user?id={ADMIN_ID}")]
        ])
        await update.message.reply_text(
            "📞 Admin-এর সাথে যোগাযোগ করতে নিচের বাটনে ক্লিক করুন।",
            reply_markup=keyboard
        )
        return ConversationHandler.END

    if text == "⚙️ Admin Panel":
        if update.effective_user.id == ADMIN_ID:
            # সরাসরি ড্যাশবোর্ড দেখাবে
            return await show_dashboard(update, context)
        else:
            await update.message.reply_text("❌ Unauthorized.")
            return ConversationHandler.END
            
    if text == "🔙 Back to Main Menu":
        return await start(update, context)
    
    return ConversationHandler.END
    
async def show_dashboard(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    global total_deposits, total_sales, balances, sales_count_per_category, user_info, transaction_log, dashboard_message

    total_balance = sum(balances.values())
    total_users_count = len(user_info)
    
    stock_info = ""
    for main_cat, sub_cats in categories.items():
        # Markdown থেকে HTML-এ পরিবর্তন: **Main_Cat** এর বদলে <b>Main_Cat</b>
        stock_info += f"  - <b>{main_cat}</b>\n" 
        for sub_cat in sub_cats:
            count = count_items(main_cat, sub_cat)
            stock_info += f"  - {sub_cat}: {count} items\n"

    sorted_sales = sorted(sales_count_per_category.items(), key=lambda item: item[1], reverse=True)
    top_selling_info = ""
    for sub_cat, count in sorted_sales:
        top_selling_info += f"  - {sub_cat}: {count} sales\n"

    recent_transactions = ""
    last_5_transactions = transaction_log[-5:]
    if last_5_transactions:
        for trans in reversed(last_5_transactions):
            trans_type, user_id, amount, timestamp = trans
            date_str = time.strftime('%H:%M %b %d', time.localtime(timestamp))
            user_data = user_info.get(user_id, {})
            
            # 🔥 মূল সংশোধন (লাইন 348-এর আশেপাশে): 'NoneType' ত্রুটির সমাধান
            # এটি নিশ্চিত করে যে user_data.get("username") যদি None হয়, তবে 'N/A' ব্যবহার করা হবে, 
            # যাতে .replace() কল করার আগে এটি একটি স্ট্রিং হয়।
            username_raw = user_data.get("username")
            username = (username_raw if username_raw is not None else "N/A").replace('_', '\\_') 
            
            # অথবা আরও সংক্ষিপ্তভাবে:
            # username = (user_data.get("username") or "N/A").replace('_', '\\_') 
            
            if trans_type == 'deposit':
                # HTML ব্যবহার করে কোড ব্লক <code></code>
                recent_transactions += f"  - 💸 ডিপোজিট: {amount} (ব্যবহারকারী: <code>@{username}</code>) at {date_str}\n" 
            elif trans_type == 'sale':
                recent_transactions += f"  - 🛒 বিক্রয়: {amount} (ব্যবহারকারী: <code>@{username}</code>) at {date_str}\n"
    else:
        recent_transactions = "  - কোনো সাম্প্রতিক লেনদেন নেই।\n"
    
    daily_deposits, daily_sales = get_report_summary(transaction_log, 1)
    weekly_deposits, weekly_sales = get_report_summary(transaction_log, 7)
    monthly_deposits, monthly_sales = get_report_summary(transaction_log, 30)

    # HTML ফরম্যাটিং ব্যবহার করে ড্যাশবোর্ড টেক্সট
    dashboard_text = (
        f"📝 <b>ড্যাশবোর্ড মেসেজ:</b>\n"
        f"<i>{dashboard_message}</i>\n" # *Italic* এর বদলে <i>Italic</i>
        "---------------------------\n"
        "📊 <b>ড্যাশবোর্ড সামারি</b>\n"
        f"👥 <b>মোট ব্যবহারকারী:</b> {total_users_count}\n"
        f"💰 <b>মোট ব্যালেন্স:</b> {sum(balances.values())} টাকা\n"
        f"🛒 <b>মোট বিক্রয়:</b> {total_sales} টাকা\n"
        f"💸 <b>মোট ডিপোজিট:</b> {total_deposits} টাকা\n"
        "---------------------------\n"
        "📈 <b>দৈনিক/সাপ্তাহিক/মাসিক রিপোর্ট</b>\n"
        f"<b>গত ২৪ ঘণ্টা:</b>\n"
        f"  - ডিপোজিট: {daily_deposits} টাকা\n"
        f"  - বিক্রয়: {daily_sales} টাকা\n"
        f"<b>গত ৭ দিন:</b>\n"
        f"  - ডিপোজিট: {weekly_deposits} টাকা\n"
        f"  - বিক্রয়: {weekly_sales} টাকা\n"
        f"<b>গত ৩০ দিন:</b>\n"
        f"  - ডিপোজিট: {monthly_deposits} টাকা\n"
        f"  - বিক্রয়: {monthly_sales} টাকা\n"
        "---------------------------\n"
        "📦 <b>বর্তমান স্টক তথ্য:</b>\n"
        f"{stock_info or '  - কোনো ক্যাটাগরি পাওয়া যায়নি।'}\n"
        "---------------------------\n"
        "📈 <b>সর্বাধিক বিক্রীত ক্যাটাগরি:</b>\n"
        f"{top_selling_info or '  - এখনো কোনো বিক্রয় হয়নি।'}\n"
        "---------------------------\n"
        "📜 <b>সর্বশেষ লেনদেন:</b>\n"
        f"{recent_transactions}\n"
    )
    
    keyboard = [
        [KeyboardButton("🔄 ড্যাশবোর্ড রিফ্রেশ করুন"), KeyboardButton("👥 ব্যবহারকারীর প্রোফাইল")],
        [KeyboardButton("📂 Manage Categories"), KeyboardButton("💰 Edit Price")],
        [KeyboardButton("✏️ Edit User Balance"), KeyboardButton("📢 Send Notice")],
        [KeyboardButton("💳 Edit Payment Info"), KeyboardButton("💳 Manage Payment Categories")],
        [KeyboardButton("🔙 Back to Main Menu")]
    ]
    
    # *** মূল পরিবর্তন: parse_mode='Markdown' থেকে parse_mode='HTML' ***
    await update.message.reply_text(dashboard_text, reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True), parse_mode='HTML')
    return ADMIN_PANEL

async def handle_dashboard_refresh(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    if update.message.text == "🔄 ড্যাশবোর্ড রিফ্রেশ করুন":
        return await show_dashboard(update, context)
    return await back_to_admin_panel_handler(update, context)

async def view_user_profile(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
    await update.message.reply_text("✍️ যে ব্যবহারকারীর প্রোফাইল দেখতে চান তার username বা User ID লিখুন:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    return SEARCH_USER_PROFILE

async def search_and_show_user_profile(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    search_term = update.message.text.strip().lstrip('@')
    
    if search_term == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    found_user_id = None
    
    # Try to search by User ID first if the input is a number
    if search_term.isdigit():
        search_id = int(search_term)
        if search_id in user_info:
            found_user_id = search_id
    
    # If not found by ID, try to search by username
    if not found_user_id:
        for user_id, info in user_info.items():
            # এখানে info.get("username") যদি None হয়, তবে .lower() কল না করেই লজিক কাজ করবে
            if info.get("username") and info["username"].lower() == search_term.lower():
                found_user_id = user_id
                break

    if not found_user_id:
        await update.message.reply_text("❌ এই ইউজারনেম বা ইউজার আইডি এর কোনো ব্যবহারকারী পাওয়া যায়নি।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return SEARCH_USER_PROFILE

    user_transactions = get_user_transactions(found_user_id, transaction_log)
    
    balance = balances.get(found_user_id, 0)
    deposits = user_deposits.get(found_user_id, 0)
    sales = user_sales.get(found_user_id, 0)

    # Calculate time-based reports
    daily_deposits, daily_sales = get_report_summary(user_transactions, 1)
    weekly_deposits, weekly_sales = get_report_summary(user_transactions, 7)
    monthly_deposits, monthly_sales = get_report_summary(user_transactions, 30)
    yearly_deposits, yearly_sales = get_report_summary(user_transactions, 365)
    
    user_data = user_info.get(found_user_id, {})
    full_name = user_data.get("first_name", "") + (f" {user_data['last_name']}" if user_data.get("last_name") else "")

    # 🔥 মূল সংশোধন: NoneType ত্রুটির সমাধান
    username_raw = user_data.get('username')
    # username_raw যদি None হয়, তবে 'N/A' ব্যবহার হবে। তারপর HTML special characters escape হবে।
    username_safe = (username_raw or 'N/A').replace('<', '&lt;').replace('>', '&gt;')
    
    profile_text = (
        f"👤 <b>ব্যবহারকারী প্রোফাইল:</b>\n"
        f"নাম: {full_name}\n"
        f"Username: <code>@{username_safe}</code>\n" # <code> tags are safe for User ID/username
        f"ID: <code>{found_user_id}</code>\n"
        "---------------------------\n"
        f"💰 <b>বর্তমান ব্যালেন্স:</b> {balance} টাকা\n"
        f"💸 <b>মোট ডিপোজিট:</b> {deposits} টাকা\n"
        f"🛒 <b>মোট খরচ:</b> {sales} টাকা\n"
        "---------------------------\n"
        "📈 <b>লেনদেনের রিপোর্ট</b>\n"
        f"<b>গত ২৪ ঘণ্টা:</b>\n"
        f"  - ডিপোজিট: {daily_deposits} টাকা\n"
        f"  - খরচ: {daily_sales} টাকা\n"
        f"<b>গত ৭ দিন:</b>\n"
        f"  - ডিপোজিট: {weekly_deposits} টাকা\n"
        f"  - খরচ: {weekly_sales} টাকা\n"
        f"<b>গত ৩০ দিন:</b>\n"
        f"  - ডিপোজিট: {monthly_deposits} টাকা\n"
        f"  - খরচ: {monthly_sales} টাকা\n"
        f"<b>গত ১ বছর:</b>\n"
        f"  - ডিপোজিট: {yearly_deposits} টাকা\n"
        f"  - খরচ: {yearly_sales} টাকা\n"
    )
    
    # *** মূল পরিবর্তন: parse_mode='HTML' ব্যবহার করা হয়েছে ***
    await update.message.reply_text(profile_text, parse_mode='HTML', reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    
    return SEARCH_USER_PROFILE

# =========================================================
# ================ NEW: BALANCE EDIT HANDLERS ===============
# =========================================================

async def edit_user_balance_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Starts the process of editing a user's balance."""
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
    
    await update.message.reply_text(
        "✏️ যে ব্যবহারকারীর ব্যালেন্স পরিবর্তন করতে চান তার username বা User ID লিখুন:",
        reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True)
    )
    return SEARCH_USER_FOR_BALANCE

async def search_user_for_balance(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Searches for the user and asks for the balance action."""
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    search_term = update.message.text.strip().lstrip('@')
    
    if search_term == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    found_user_id = None
    
    if search_term.isdigit():
        search_id = int(search_term)
        if search_id in user_info:
            found_user_id = search_id
    
    if not found_user_id:
        for user_id, info in user_info.items():
            if info.get("username") and info["username"].lower() == search_term.lower():
                found_user_id = user_id
                break

    if not found_user_id:
        await update.message.reply_text(
            "❌ এই ইউজারনেম বা ইউজার আইডি এর কোনো ব্যবহারকারী পাওয়া যায়নি। আবার চেষ্টা করুন।",
            reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True)
        )
        return SEARCH_USER_FOR_BALANCE

    # Store found user ID for the next steps
    context.user_data['edit_balance_user_id'] = found_user_id
    
    user_data = user_info.get(found_user_id, {})
    username = user_data.get('username', 'N/A')
    current_balance = balances.get(found_user_id, 0)

    keyboard = [
        [KeyboardButton("➕ Add Balance"), KeyboardButton("➖ Remove Balance")],
        [KeyboardButton("✍️ Set New Balance")],
        [KeyboardButton("🔙 অ্যাডমিন প্যানেল")]
    ]
    
    await update.message.reply_text(
        f"👤 ব্যবহারকারী: @{username} (ID: {found_user_id})\n"
        f"💰 বর্তমান ব্যালেন্স: {current_balance} টাকা\n\n"
        f"আপনি কী করতে চান?",
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    )
    
    return BALANCE_EDIT_ACTION

async def balance_edit_action_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Handles the admin's choice of action (Add, Remove, Set)."""
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    action_text = update.message.text
    
    if action_text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    if action_text not in ["➕ Add Balance", "➖ Remove Balance", "✍️ Set New Balance"]:
        await update.message.reply_text("❌ অনুগ্রহ করে নিচের বাটন থেকে একটি অপশন বেছে নিন।")
        return BALANCE_EDIT_ACTION

    context.user_data['balance_edit_action'] = action_text
    
    if action_text == "➕ Add Balance":
        prompt = "✍️ কত টাকা যোগ করতে চান তা লিখুন:"
    elif action_text == "➖ Remove Balance":
        prompt = "✍️ কত টাকা সরাতে চান তা লিখুন:"
    else: # "✍️ Set New Balance"
        prompt = "✍️ নতুন ব্যালেন্স কত হবে তা লিখুন:"
        
    await update.message.reply_text(
        prompt,
        reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True)
    )
    
    return RECEIVE_BALANCE_EDIT_AMOUNT

async def receive_balance_edit_amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Receives the amount and performs the balance update."""
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    amount_str = update.message.text.strip()
    
    if amount_str == "🔙 অ্যাডমিন প্যানেল":
        context.user_data.pop('edit_balance_user_id', None)
        context.user_data.pop('balance_edit_action', None)
        return await back_to_admin_panel_handler(update, context)

    if not amount_str.isdigit() or float(amount_str) < 0:
        await update.message.reply_text(
            "❌ অনুগ্রহ করে একটি ধনাত্মক সংখ্যা লিখুন।",
            reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True)
        )
        return RECEIVE_BALANCE_EDIT_AMOUNT
        
    amount = float(amount_str)
    user_id = context.user_data.get('edit_balance_user_id')
    action = context.user_data.get('balance_edit_action')

    if not user_id or not action:
        await update.message.reply_text("❌ একটি ত্রুটি ঘটেছে। অনুগ্রহ করে আবার শুরু করুন।")
        return await back_to_admin_panel_handler(update, context)

    old_balance = balances.get(user_id, 0)
    new_balance = 0
    
    if action == "➕ Add Balance":
        new_balance = old_balance + amount
        balances[user_id] = new_balance
        user_message = f"✅ অ্যাডমিন আপনার ব্যালেন্সে {amount} টাকা যোগ করেছে।\nআপনার নতুন ব্যালেন্স: {new_balance} টাকা।"
        admin_message = f"✅ ব্যবহারকারীর ব্যালেন্সে {amount} টাকা যোগ করা হয়েছে।\nনতুন ব্যালেন্স: {new_balance} টাকা।"

    elif action == "➖ Remove Balance":
        new_balance = old_balance - amount
        if new_balance < 0:
            new_balance = 0 # Balance can't be negative
        balances[user_id] = new_balance
        user_message = f"✅ অ্যাডমিন আপনার ব্যালেন্স থেকে {amount} টাকা সরিয়ে নিয়েছে।\nআপনার নতুন ব্যালেন্স: {new_balance} টাকা।"
        admin_message = f"✅ ব্যবহারকারীর ব্যালেন্স থেকে {amount} টাকা সরানো হয়েছে।\nনতুন ব্যালেন্স: {new_balance} টাকা।"

    elif action == "✍️ Set New Balance":
        new_balance = amount
        balances[user_id] = new_balance
        user_message = f"✅ অ্যাডমিন আপনার নতুন ব্যালেন্স {new_balance} টাকা সেট করেছে।"
        admin_message = f"✅ ব্যবহারকারীর নতুন ব্যালেন্স {new_balance} টাকা সেট করা হয়েছে।"

    save_user_data()

    # Notify the user
    try:
        await context.bot.send_message(chat_id=user_id, text=user_message)
    except Exception as e:
        logger.error(f"Failed to notify user {user_id} about balance change: {e}")
        admin_message += "\n⚠️ ব্যবহারকারীকে নোটিশ পাঠানো যায়নি।"

    # Confirm to admin and go back to panel
    await update.message.reply_text(admin_message)
    
    context.user_data.pop('edit_balance_user_id', None)
    context.user_data.pop('balance_edit_action', None)
    
    return await show_dashboard(update, context)


async def deposit_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    amount_str = update.message.text.strip()

    if amount_str == "🔙 Back to Main Menu":
        return await start(update, context)

    if not amount_str.isdigit():
        await update.message.reply_text("❌ অনুগ্রহ করে শুধু সংখ্যা লিখুন।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
        return DEPOSIT

    amount = int(amount_str)
    
    context.user_data["deposit_amount"] = amount
    
    keyboard = [
        [KeyboardButton("🔙 Back to Main Menu")]
    ]
    
    await update.message.reply_text(
        f"💳 এখন অনুগ্রহ করে {amount} টাকা পেমেন্ট করুন:\n{payment_info}\n\n"
        f"📸 পেমেন্টের পর স্ক্রিনশট পাঠান।",
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    )
    return GET_DEPOSIT_AMOUNT

async def receive_deposit_screenshot(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text == "🔙 Back to Main Menu":
        return await start(update, context)

    if not update.message.photo:
        await update.message.reply_text("❌ অনুগ্রহ করে শুধু একটি ছবি পাঠান।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
        return GET_DEPOSIT_AMOUNT
    
    user = update.effective_user
    deposit_amount = context.user_data.get("deposit_amount", 0)
    
    username = user.username if user.username else 'N/A'
    caption = (
        f"🔔 **নতুন ডিপোজিট রিকোয়েস্ট!** 🔔\n"
        f"User: @{username}\n"
        f"Amount: {deposit_amount}\n"
        f"UserID: {user.id}"
    )

    keyboard = InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Confirm Deposit", callback_data=f"deposit_confirm:{user.id}:{deposit_amount}"),
         InlineKeyboardButton("❌ Cancel Deposit", callback_data=f"deposit_cancel:{user.id}")]
    ])

    await context.bot.send_photo(chat_id=ADMIN_ID, photo=update.message.photo[-1].file_id, caption=caption, reply_markup=keyboard)
    
    await update.message.reply_text("✅ আপনার ডিপোজিট রিকোয়েস্ট Admin-কে পাঠানো হয়েছে। Admin নিশ্চিত করার পর আপনার ব্যালেন্স যোগ হবে।")
    
    context.user_data.clear()
    return ConversationHandler.END

async def back_to_admin_panel_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    # এখন Admin Panel-এ ফেরার মানেই ড্যাশবোর্ড দেখানো
    return await show_dashboard(update, context)

async def admin_panel_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    text = update.message.text
    
    if text == "🔙 Back to Main Menu":
        return await start(update, context)
    
    # NEW: ড্যাশবোর্ড রিফ্রেশ হ্যান্ডেল করা হচ্ছে
    if text == "🔄 ড্যাশবোর্ড রিফ্রেশ করুন":
        # ড্যাশবোর্ড রিফ্রেশ করতে show_dashboard কল হবে
        return await show_dashboard(update, context)
        
    # NOTE: OLD 'if text == "📊 Dashboard":' ব্লকটি এখন আর এখানে নেই।
    
    if text == "👥 ব্যবহারকারীর প্রোফাইল":
        return await view_user_profile(update, context)

    # --- NEW: Handle Edit User Balance ---
    if text == "✏️ Edit User Balance":
        return await edit_user_balance_start(update, context)
        
    if text == "📢 Send Notice":
        await update.message.reply_text("✍️ যে নোটিশটি পাঠাতে চান তা লিখুন:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return SEND_NOTICE
        
    if text == "📂 Manage Categories":
        keyboard = []
        for cat in categories.keys():
            stock_count = get_total_stock(cat)
            keyboard.append([KeyboardButton(f"{cat} ({stock_count})")])
        
        keyboard.append([KeyboardButton("➕ Add Main Category")])
        keyboard.append([KeyboardButton("➖ Remove Main Category")])
        keyboard.append([KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        await update.message.reply_text("⚙️ Manage Main Categories:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return MANAGE_CATEGORY

    if text == "💰 Edit Price":
        keyboard = [[KeyboardButton(cat)] for cat in categories.keys()]
        keyboard.append([KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        await update.message.reply_text("✍️ কোন Category-র মূল্য পরিবর্তন করবেন?", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return EDIT_PRICE_MAIN

    if text == "💳 Edit Payment Info":
        await update.message.reply_text("✍️ নতুন পেমেন্ট তথ্য পাঠান:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return EDIT_PAYMENT
        
    if text == "💳 Manage Payment Categories":
        return await manage_payment_categories_handler(update, context)
         
    return ADMIN_PANEL

async def manage_payment_categories_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    keyboard_inline = []
    if not categories:
        await update.message.reply_text("⚠️ কোনো ক্যাটাগরি নেই।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return ADMIN_PANEL

    for main_cat in categories:
        current_status = "Manual Payment 💳" if main_cat in MANUAL_DELIVERY_CATEGORIES else "Balance Payment 💰"
        button_text = "Switch to Balance" if main_cat in MANUAL_DELIVERY_CATEGORIES else "Switch to Manual"
        callback_data = f"toggle_payment:{main_cat}"
        
        keyboard_inline.append([
            InlineKeyboardButton(f"{main_cat} ({current_status})", callback_data="ignore"),
            InlineKeyboardButton(button_text, callback_data=callback_data)
        ])
    
    reply_markup_inline = InlineKeyboardMarkup(keyboard_inline)
    
    # Send the main message with inline buttons
    await update.message.reply_text(
        "⚡️ **পেমেন্ট ক্যাটাগরি নিয়ন্ত্রণ**\n\n"
        "নিচের তালিকা থেকে প্রতিটি ক্যাটাগরির জন্য পেমেন্ট পদ্ধতি পরিবর্তন করতে পারেন।\n"
        "**Balance Payment** মানে ব্যবহারকারী তার ব্যালেন্স থেকে কিনতে পারবে।\n"
        "**Manual Payment** মানে ব্যবহারকারীকে সরাসরি পেমেন্ট করে স্ক্রিনশট পাঠাতে হবে।",
        reply_markup=reply_markup_inline,
        parse_mode='Markdown'
    )

    # Send the second message with the back button
    reply_markup_text = ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True)
    await update.message.reply_text(
        "🔙 অ্যাডমিন প্যানেল-এ ফিরে যেতে নিচের বাটনটি চাপুন।",
        reply_markup=reply_markup_text
    )

    return MANAGE_PAYMENT_CATEGORIES

async def toggle_payment_method(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data
    
    if update.effective_user.id != ADMIN_ID:
        await query.edit_message_caption("❌ Unauthorized.", reply_markup=None)
        return
        
    if data.startswith("toggle_payment:"):
        _, cat_name = data.split(":", 1)
        
        if cat_name in MANUAL_DELIVERY_CATEGORIES:
            MANUAL_DELIVERY_CATEGORIES.remove(cat_name)
        else:
            MANUAL_DELIVERY_CATEGORIES.append(cat_name)
        
        save_user_data()
        
        # Rebuild the keyboard with the updated status
        keyboard = []
        for main_cat in categories:
            current_status = "Manual Payment 💳" if main_cat in MANUAL_DELIVERY_CATEGORIES else "Balance Payment 💰"
            button_text = "Switch to Balance" if main_cat in MANUAL_DELIVERY_CATEGORIES else "Switch to Manual"
            callback_data = f"toggle_payment:{main_cat}"
            keyboard.append([
                InlineKeyboardButton(f"{main_cat} ({current_status})", callback_data="ignore"),
                InlineKeyboardButton(button_text, callback_data=callback_data)
            ])
            
        reply_markup = InlineKeyboardMarkup(keyboard)
        await query.edit_message_text(
            "⚡️ **পেমেন্ট ক্যাটাগরি নিয়ন্ত্রণ**\n\n"
            "নিচের তালিকা থেকে প্রতিটি ক্যাটাগরির জন্য পেমেন্ট পদ্ধতি পরিবর্তন করতে পারেন।\n"
            "**Balance Payment** মানে ব্যবহারকারী তার ব্যালেন্স থেকে কিনতে পারবে।\n"
            "**Manual Payment** মানে ব্যবহারকারীকে সরাসরি পেমেন্ট করে স্ক্রিনশট পাঠাতে হবে।",
            reply_markup=reply_markup,
            parse_mode='Markdown'
        )

    # In case of "ignore" data, do nothing
    return

async def send_notice_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
    
    notice_text = update.message.text
    
    notice_count = 0
    failed_users = []
    
    # Get all users who have started the bot, excluding the admin
    users_to_notify = [uid for uid in user_info.keys() if uid != ADMIN_ID]
    
    if not users_to_notify:
        await update.message.reply_text("⚠️ কোনো ব্যবহারকারীকে নোটিশ পাঠানো হয়নি। সম্ভবত অন্য কোনো ব্যবহারকারী এখনও বট শুরু করেনি।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return await back_to_admin_panel_handler(update, context)

    for user_id in users_to_notify:
        try:
            await context.bot.send_message(chat_id=user_id, text=f"📢 **নোটিশ:**\n\n{notice_text}", parse_mode='Markdown')
            notice_count += 1
        except Exception:
            # Handle cases where the user has blocked the bot
            failed_users.append(user_id)
            pass

    await update.message.reply_text(f"✅ নোটিশ পাঠানো হয়েছে।\n\nসফল: {notice_count} জন ব্যবহারকারীকে\nব্যর্থ: {len(failed_users)} জন ব্যবহারকারীকে", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    return await back_to_admin_panel_handler(update, context)

async def back_to_manage_main_categories_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """A helper function to handle the return to the main category management view."""
    if "active_main_cat" in context.user_data:
        del context.user_data["active_main_cat"]
    if "active_sub_cat" in context.user_data:
        del context.user_data["active_sub_cat"]
    
    return await manage_category_handler(update, context)

async def manage_category_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    text = update.message.text
    
    if text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    if text == "➕ Add Main Category":
        await update.message.reply_text("✍️ নতুন প্রধান Category-র নাম পাঠান:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return ADD_MAIN_CAT

    if text == "➖ Remove Main Category":
        keyboard = []
        for cat in categories.keys():
            stock_count = get_total_stock(cat)
            keyboard.append([KeyboardButton(f"{cat} ({stock_count})")])
        
        keyboard.append([KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        await update.message.reply_text("➖ কোন প্রধান Category remove করবেন?", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return REMOVE_MAIN_CAT
    
    original_cat = text.split(" (")[0]
    
    if original_cat in categories:
        context.user_data["active_main_cat"] = original_cat
        keyboard = []
        for sub_cat in categories[original_cat]:
            stock_count = count_items(original_cat, sub_cat)
            keyboard.append([KeyboardButton(f"{sub_cat} ({stock_count})")])
        
        keyboard.append([KeyboardButton("➕ Add Sub Category")])
        keyboard.append([KeyboardButton("➖ Remove Sub Category")])
        keyboard.append([KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        
        await update.message.reply_text(f"⚙️ Manage Sub Categories for **{original_cat}**:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return MANAGE_SUB_CATEGORY
    
    return MANAGE_CATEGORY

async def add_main_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    new_cat = update.message.text.strip()
    
    if new_cat == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
        
    if new_cat in categories:
        await update.message.reply_text("⚠️ এই প্রধান ক্যাটাগরি আগে থেকেই আছে।")
    else:
        categories[new_cat] = []
        save_user_data()
        await update.message.reply_text(f"✅ প্রধান ক্যাটাগরি '{new_cat}' যোগ হয়েছে।")
    
    return await manage_category_handler(update, context)

async def remove_main_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    cat_to_remove = update.message.text.split(" (")[0].strip()
    
    if cat_to_remove == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
        
    if cat_to_remove in categories:
        # Remove associated Excel files
        if cat_to_remove in categories:
            for sub_cat in categories[cat_to_remove]:
                path = get_excel_path(cat_to_remove, sub_cat)
                if os.path.exists(path):
                    os.remove(path)
                if cat_to_remove in prices and sub_cat in prices[cat_to_remove]:
                    del prices[cat_to_remove][sub_cat]
        if cat_to_remove in prices:
            del prices[cat_to_remove]
        
        del categories[cat_to_remove]
        # Also remove from MANUAL_DELIVERY_CATEGORIES if it exists there
        if cat_to_remove in MANUAL_DELIVERY_CATEGORIES:
            MANUAL_DELIVERY_CATEGORIES.remove(cat_to_remove)
            
        save_user_data()
        await update.message.reply_text(f"✅ প্রধান ক্যাটাগরি '{cat_to_remove}' remove হয়েছে।")
    else:
        await update.message.reply_text("⚠️ এই ক্যাটাগরি পাওয়া যায়নি।")
    
    return await manage_category_handler(update, context)
    
async def manage_sub_category_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    text = update.message.text
    main_cat = context.user_data.get("active_main_cat")
    
    if text == "🔙 Manage Main Categories":
        if "active_main_cat" in context.user_data:
            del context.user_data["active_main_cat"]
        return await manage_category_handler(update, context)
    
    if text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    if text == "➕ Add Sub Category":
        await update.message.reply_text("✍️ নতুন সাব-ক্যাটাগরির নাম পাঠান:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return ADD_SUB_CAT
        
    if text == "➖ Remove Sub Category":
        keyboard = [[KeyboardButton(sub_cat)] for sub_cat in categories.get(main_cat, [])]
        keyboard.append([KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        await update.message.reply_text("➖ কোন সাব-ক্যাটাগরি remove করবেন?", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return REMOVE_SUB_CAT
    
    original_sub_cat = text.split(" (")[0]
    
    if original_sub_cat in categories.get(main_cat, []):
        context.user_data["active_sub_cat"] = original_sub_cat
        count = count_items(main_cat, original_sub_cat)
        keyboard = [
            [KeyboardButton("➕ Add Items")],
            [KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")]
        ]
        await update.message.reply_text(
            f"⚙️ Sub-Category: {original_sub_cat}\n📦 Items in stock: {count}\n\nআইটেম যোগ করতে `➕ Add Items` চাপুন।",
            reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
        )
        return ADD_ITEMS
        
    return MANAGE_SUB_CATEGORY

async def add_sub_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    new_sub_cat = update.message.text.strip()
    main_cat = context.user_data.get("active_main_cat")
    
    if new_sub_cat == "🔙 Manage Main Categories":
        await back_to_manage_main_categories_handler(update, context)
        return MANAGE_CATEGORY

    if new_sub_cat == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    if main_cat and new_sub_cat not in categories[main_cat]:
        categories[main_cat].append(new_sub_cat)
        ensure_excel(main_cat, new_sub_cat)
        save_user_data()
        await update.message.reply_text(f"✅ সাব-ক্যাটাগরি '{new_sub_cat}' যোগ হয়েছে।")
    else:
        await update.message.reply_text("⚠️ এই সাব-ক্যাটাগরি আগে থেকেই আছে অথবা কোনো প্রধান ক্যাটাগরি নির্বাচন করা হয়নি।")
    
    # After adding or handling the error, display the correct MANAGE_SUB_CATEGORY keyboard
    keyboard = []
    for sub_cat in categories.get(main_cat, []):
        stock_count = count_items(main_cat, sub_cat)
        keyboard.append([KeyboardButton(f"{sub_cat} ({stock_count})")])
    
    keyboard.append([KeyboardButton("➕ Add Sub Category")])
    keyboard.append([KeyboardButton("➖ Remove Sub Category")])
    keyboard.append([KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
    
    await update.message.reply_text(
        f"⚙️ Manage Sub Categories for **{main_cat}**:", 
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    )
    
    return MANAGE_SUB_CATEGORY
    
async def remove_sub_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    sub_cat_to_remove = update.message.text.split(" (")[0].strip()
    main_cat = context.user_data.get("active_main_cat")
    
    if sub_cat_to_remove == "🔙 Manage Main Categories":
        await back_to_manage_main_categories_handler(update, context)
        return MANAGE_CATEGORY

    if sub_cat_to_remove == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
        
    if main_cat and sub_cat_to_remove in categories.get(main_cat, []):
        categories[main_cat].remove(sub_cat_to_remove)
        path = get_excel_path(main_cat, sub_cat_to_remove)
        if os.path.exists(path):
            os.remove(path)
        if main_cat in prices and sub_cat_to_remove in prices[main_cat]:
            del prices[main_cat][sub_cat_to_remove]
        save_user_data()
        await update.message.reply_text(f"✅ সাব-ক্যাটাগরি '{sub_cat_to_remove}' remove হয়েছে।")
    else:
        await update.message.reply_text("⚠️ এই সাব-ক্যাটাগরি পাওয়া যায়নি।")
    
    # After removing or handling the error, display the correct MANAGE_SUB_CATEGORY keyboard
    keyboard = []
    for sub_cat in categories.get(main_cat, []):
        stock_count = count_items(main_cat, sub_cat)
        keyboard.append([KeyboardButton(f"{sub_cat} ({stock_count})")])
    
    keyboard.append([KeyboardButton("➕ Add Sub Category")])
    keyboard.append([KeyboardButton("➖ Remove Sub Category")])
    keyboard.append([KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
    
    await update.message.reply_text(
        f"⚙️ Manage Sub Categories for **{main_cat}**:", 
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    )
    
    return MANAGE_SUB_CATEGORY

async def add_item_line(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    text = update.message.text
    
    if text == "🔙 Manage Main Categories":
        await back_to_manage_main_categories_handler(update, context)
        return MANAGE_CATEGORY

    if text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
    
    main_cat = context.user_data.get("active_main_cat")
    sub_cat = context.user_data.get("active_sub_cat")
    
    if not main_cat or not sub_cat:
        await update.message.reply_text("⚠️ কোনো ক্যাটাগরি নির্বাচন করা হয়নি।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return await manage_category_handler(update, context)
        
    if text == "➕ Add Items":
        await update.message.reply_text(f"✍️ '{sub_cat}' এর জন্য আইটেম পাঠান (প্রতি লাইনে একটি করে)।\n\nসব আইটেম পাঠানো হলে `✅ Done` চাপুন।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("✅ Done")], [KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return ADD_ITEMS
        
    if text == "✅ Done":
        count = count_items(main_cat, sub_cat)
        await update.message.reply_text(f"✅ '{sub_cat}' তে মোট {count} আইটেম আছে।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        if "active_main_cat" in context.user_data:
            del context.user_data["active_main_cat"]
        if "active_sub_cat" in context.user_data:
            del context.user_data["active_sub_cat"]
        return await back_to_admin_panel_handler(update, context)
    
    if text:
        items_to_add = text.split('\n')
        added_count = 0
        for item in items_to_add:
            item = item.strip()
            if item:
                add_item_to_excel(main_cat, sub_cat, item)
                added_count += 1
        await update.message.reply_text(f"✅ Added {added_count} item(s).", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("✅ Done")], [KeyboardButton("🔙 Manage Main Categories"), KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    
    return ADD_ITEMS

async def edit_payment_info(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    global payment_info
    new_info = update.message.text.strip()
    
    if new_info == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
        
    payment_info = new_info
    await update.message.reply_text("✅ পেমেন্ট তথ্য সফলভাবে আপডেট হয়েছে।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    return await back_to_admin_panel_handler(update, context)
    
async def edit_price_main_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    text = update.message.text.strip()
    
    if text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)
    
    original_cat = text.split(" (")[0]
        
    if original_cat in categories.keys():
        context.user_data['temp_main_cat_for_price'] = original_cat
        keyboard = [[KeyboardButton(sub_cat)] for sub_cat in categories[original_cat]]
        keyboard.append([KeyboardButton("🔙 অ্যাডমিন প্যানেল")])
        await update.message.reply_text(f"✍️ কোন সাব-ক্যাটাগরির মূল্য পরিবর্তন করবেন?", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return EDIT_PRICE_SUB
        
    await update.message.reply_text("❌ এই ক্যাটাগরি পাওয়া যায়নি।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    return EDIT_PRICE_MAIN

async def edit_price_sub_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END
        
    text = update.message.text.strip()
    main_cat = context.user_data.get('temp_main_cat_for_price')
    
    if text == "🔙 অ্যাডমিন প্যানেল":
        if 'temp_main_cat_for_price' in context.user_data:
            del context.user_data['temp_main_cat_for_price']
        return await back_to_admin_panel_handler(update, context)
        
    if main_cat and text in categories.get(main_cat, []):
        context.user_data['temp_sub_cat_for_price'] = text
        current_price = prices.get(main_cat, {}).get(text, "সেট করা হয়নি")
        await update.message.reply_text(f"✍️ '{text}' এর বর্তমান মূল্য: {current_price}\nনতুন মূল্য লিখুন:", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return RECEIVE_NEW_PRICE
        
    await update.message.reply_text("❌ এই সাব-ক্যাটাগরি পাওয়া যায়নি।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
    return EDIT_PRICE_SUB

async def receive_price(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return ConversationHandler.END

    price_text = update.message.text.strip()
    
    if price_text == "🔙 অ্যাডমিন প্যানেল":
        if 'temp_main_cat_for_price' in context.user_data:
            del context.user_data['temp_main_cat_for_price']
        if 'temp_sub_cat_for_price' in context.user_data:
            del context.user_data['temp_sub_cat_for_price']
        return await back_to_admin_panel_handler(update, context)

    try:
        new_price = float(price_text)
        main_cat = context.user_data.get('temp_main_cat_for_price')
        sub_cat = context.user_data.get('temp_sub_cat_for_price')
        if main_cat and sub_cat:
            if main_cat not in prices:
                prices[main_cat] = {}
            prices[main_cat][sub_cat] = new_price
            save_user_data()
            await update.message.reply_text(f"✅ '{sub_cat}' এর মূল্য এখন {new_price}।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
            if 'temp_main_cat_for_price' in context.user_data:
                del context.user_data['temp_main_cat_for_price']
            if 'temp_sub_cat_for_price' in context.user_data:
                del context.user_data['temp_sub_cat_for_price']
            return await back_to_admin_panel_handler(update, context)
    except ValueError:
        await update.message.reply_text("❌ মূল্য শুধুমাত্র সংখ্যায় লিখুন।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 অ্যাডমিন প্যানেল")]], resize_keyboard=True))
        return RECEIVE_NEW_PRICE
        
    return await back_to_admin_panel_handler(update, context)

async def back_to_categories_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END

    keyboard = []
    for cat in categories.keys():
        keyboard.append([KeyboardButton(cat)])
    
    keyboard.append([KeyboardButton("🔙 Back to Main Menu")])
    await update.message.reply_text("🛒 ক্যাটাগরি বেছে নিন:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
    return BUY_MENU
    
async def back_to_subcategories_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END
    
    main_cat = context.user_data.get('temp_main_cat_for_buy')
    if not main_cat:
        return await back_to_categories_handler(update, context)
        
    keyboard = []
    for sub_cat in categories[main_cat]:
        stock_count = count_items(main_cat, sub_cat)
        keyboard.append([KeyboardButton(f"{sub_cat} ({stock_count})")])
        
    keyboard.append([KeyboardButton("🔙 Back to Categories")])
    keyboard.append([KeyboardButton("🔙 Back to Main Menu")])

    await update.message.reply_text(f"🛒 **{main_cat}** এর সাব-ক্যাটাগরি বেছে নিন:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
    return BUY_SUB_MENU

async def user_choose_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END
        
    text = update.message.text.strip()
    original_cat = text.split(" (")[0]
    
    if original_cat == "🔙 Back to Main Menu":
        return await start(update, context)

    context.user_data.clear()

    if original_cat in categories.keys():
        context.user_data["temp_main_cat_for_buy"] = original_cat
        
        keyboard = []
        for sub_cat in categories[original_cat]:
            stock_count = count_items(original_cat, sub_cat)
            keyboard.append([KeyboardButton(f"{sub_cat} ({stock_count})")])
        
        keyboard.append([KeyboardButton("🔙 Back to Categories")])
        keyboard.append([KeyboardButton("🔙 Back to Main Menu")])
        
        await update.message.reply_text(f"🛒 **{original_cat}** এর সাব-ক্যাটাগরি বেছে নিন:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return BUY_SUB_MENU

    await update.message.reply_text("❌ এই ক্যাটাগরি পাওয়া যায়নি।")
    return BUY_MENU

async def user_choose_subcategory(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END
        
    text = update.message.text.strip()
    original_sub_cat = text.split(" (")[0]
    main_cat = context.user_data.get('temp_main_cat_for_buy')
    
    if text == "🔙 Back to Categories":
        return await back_to_categories_handler(update, context)
    if text == "🔙 Back to Main Menu":
        return await start(update, context)
        
    if main_cat and original_sub_cat in categories.get(main_cat, []):
        context.user_data["order"] = {"main_cat": main_cat, "sub_cat": original_sub_cat}
        
        price = prices.get(main_cat, {}).get(original_sub_cat, "মূল্য এখনো সেট করা হয়নি।")
        
        keyboard = [
            [KeyboardButton("🔙 Back to Sub-Categories")],
            [KeyboardButton("🔙 Back to Main Menu")]
        ]
        
        await update.message.reply_text(f"✅ আপনি **{original_sub_cat}** সিলেক্ট করেছেন।\n💰 **প্রতিটির দাম:** {price} টাকা\n\n✍️ অনুগ্রহ করে আপনি কতগুলি চান তা সংখ্যায় লিখুন:", reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True))
        return GET_QUANTITY

    await update.message.reply_text("❌ এই সাব-ক্যাটাগরি পাওয়া যায়নি।")
    return BUY_SUB_MENU


async def receive_quantity(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END

    qty = update.message.text.strip()

    if qty == "🔙 Back to Sub-Categories":
        return await back_to_subcategories_handler(update, context)
    if qty == "🔙 Back to Main Menu":
        return await start(update, context)
    
    if not qty.isdigit():
        await update.message.reply_text("❌ অনুগ্রহ করে শুধু সংখ্যা লিখুন।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Sub-Categories"), KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
        return GET_QUANTITY
        
    qty = int(qty)
    order = context.user_data.get("order", {})
    order["qty"] = qty
    
    main_cat = order['main_cat']
    sub_cat = order['sub_cat']
    
    price_per_item = prices.get(main_cat, {}).get(sub_cat, 0)
    total_price = price_per_item * qty
    
    order["price"] = total_price
    context.user_data["order"] = order

    is_manual = main_cat in MANUAL_DELIVERY_CATEGORIES

    if not is_manual:
        user_id = update.effective_user.id
        current_balance = balances.get(user_id, 0)
    
        if current_balance >= total_price:
            keyboard = [
                [KeyboardButton("✅ Confirm Purchase")],
                [KeyboardButton("🔙 Back to Sub-Categories")],
                [KeyboardButton("🔙 Back to Main Menu")]
            ]
            await update.message.reply_text(
                f"✅ অর্ডার তৈরি হয়েছে।\n"
                f"ক্যাটাগরি: {order['sub_cat']}\n"
                f"পরিমাণ: {order['qty']}\n"
                f"মোট দাম: {total_price} টাকা\n"
                f"আপনার বর্তমান ব্যালেন্স: {current_balance} টাকা\n\n"
                f"আপনার ব্যালেন্স থেকে পেমেন্ট করতে `✅ Confirm Purchase` চাপুন।",
                reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
            )
        else:
            await update.message.reply_text(f"❌ আপনার যথেষ্ট ব্যালেন্স নেই। আপনার প্রয়োজন {total_price} টাকা কিন্তু ব্যালেন্স আছে {current_balance} টাকা।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Sub-Categories"), KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
            return BUY_SUB_MENU
            
        return WAIT_SCREENSHOT
    
    keyboard = [
        [KeyboardButton("🔙 Back to Sub-Categories")],
        [KeyboardButton("🔙 Back to Main Menu")]
    ]
    await update.message.reply_text(
        f"✅ অর্ডার তৈরি হয়েছে।\n"
        f"ক্যাটাগরি: {order['sub_cat']}\n"
        f"পরিমাণ: {order['qty']}\n"
        f"মোট দাম: {total_price} টাকা\n"
        f"⚠️ অনুগ্রহ করে পেমেন্ট করুন:\n{payment_info}\n\n"
        f"📸 পেমেন্টের পর স্ক্রিনশট পাঠান।",
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    )
    return WAIT_SCREENSHOT


async def user_send_screenshot(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await check_subscription(update, context):
        return ConversationHandler.END

    if update.message.text == "🔙 Back to Sub-Categories":
        return await back_to_subcategories_handler(update, context)
    if update.message.text == "🔙 Back to Main Menu":
        return await start(update, context)
    if update.message.text == "🔙 অ্যাডমিন প্যানেল":
        return await back_to_admin_panel_handler(update, context)

    order = context.user_data.get("order", {})
    if not order:
        return ConversationHandler.END
        
    main_cat = order['main_cat']
    sub_cat = order['sub_cat']

    if update.message.text == "✅ Confirm Purchase" and main_cat not in MANUAL_DELIVERY_CATEGORIES:
        user_id = update.effective_user.id
        total_price = order['price']
        qty = order['qty']
        current_balance = balances.get(user_id, 0)
        
        if current_balance >= total_price:
            balances[user_id] = current_balance - total_price
            
            items = pop_items_from_excel(main_cat, sub_cat, qty)
            if not items:
                balances[user_id] = balances.get(user_id, 0) + total_price 
                await update.message.reply_text("❌ যথেষ্ট আইটেম স্টক এ নেই। আপনার ব্যালেন্স রিফান্ড করা হয়েছে।", reply_markup=ReplyKeyboardRemove())
                return await start(update, context)
            
            global total_sales, sales_count_per_category, transaction_log, user_sales
            total_sales += total_price
            sales_count_per_category[sub_cat] = sales_count_per_category.get(sub_cat, 0) + qty
            transaction_log.append(('sale', user_id, total_price, time.time()))
            user_sales[user_id] = user_sales.get(user_id, 0) + total_price

            item_text = "\n".join(items)
            await context.bot.send_message(chat_id=user_id, text=f"✅ আপনার অর্ডার:\n{item_text}")

            await update.message.reply_text("✅ ব্যালেন্স ব্যবহার করে আপনার অর্ডার সফল হয়েছে!", reply_markup=ReplyKeyboardRemove())
            context.user_data.clear()
            save_user_data()
            return await start(update, context)
        else:
            await update.message.reply_text("❌ আপনার যথেষ্ট ব্যালেন্স নেই।", reply_markup=ReplyKeyboardRemove())
            return await start(update, context)

    # For manual payment via screenshot
    if not update.message.photo:
        await update.message.reply_text("❌ অনুগ্রহ করে শুধু একটি ছবি পাঠান।", reply_markup=ReplyKeyboardMarkup([[KeyboardButton("🔙 Back to Sub-Categories"), KeyboardButton("🔙 Back to Main Menu")]], resize_keyboard=True))
        return WAIT_SCREENSHOT
        
    user = update.effective_user
    username = user.username if user.username else 'N/A'
    caption = (
        f"🔔 **নতুন অর্ডার!** 🔔\n"
        f"User: @{username}\n"
        f"Category: {order['sub_cat']}\n"
        f"Quantity: {order['qty']}\n"
        f"Price: {order['price']}\n"
        f"UserID: {user.id}"
    )

    keyboard = InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Confirm Order", callback_data=f"confirm_manual:{user.id}:{order['main_cat']}:{order['sub_cat']}:{order['qty']}:{order['price']}"),
         InlineKeyboardButton("❌ Cancel Order", callback_data=f"cancel_manual:{user.id}")]
    ])
    
    await context.bot.send_photo(chat_id=ADMIN_ID, photo=update.message.photo[-1].file_id, caption=caption, reply_markup=keyboard)
    
    await update.message.reply_text("✅ Screenshot পাঠানো হয়েছে Admin-কে।")
    await start(update, context)
    context.user_data.clear()
    return ConversationHandler.END

async def admin_order_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    global total_sales, sales_count_per_category, transaction_log, user_sales, prices
    query = update.callback_query
    await query.answer()
    data = query.data
    
    if update.effective_user.id != ADMIN_ID:
        await query.edit_message_caption("❌ Unauthorized.", reply_markup=None)
        return
    
    if data.startswith("confirm:"): # Auto-delivery via balance
        parts = data.split(":")
        uid = int(parts[1])
        main_cat = parts[2]
        sub_cat = parts[3]
        qty = int(parts[4])

        items = pop_items_from_excel(main_cat, sub_cat, qty)
        
        if not items:
            await query.edit_message_caption(query.message.caption + "\n\n❌ Not enough stock.", reply_markup=None)
            return

        price_per_item = prices.get(main_cat, {}).get(sub_cat, 0)
        total_price = price_per_item * qty
        total_sales += total_price
        sales_count_per_category[sub_cat] = sales_count_per_category.get(sub_cat, 0) + qty
        transaction_log.append(('sale', uid, total_price, time.time()))
        user_sales[uid] = user_sales.get(uid, 0) + total_price
        save_user_data()

        item_text = "\n".join(items)
        await context.bot.send_message(chat_id=uid, text=f"✅ আপনার অর্ডার:\n{item_text}")

        await query.edit_message_caption(query.message.caption + "\n\n✅ Delivered (Text Sent)", reply_markup=None)
    
    elif data.startswith("confirm_manual:"):
        parts = data.split(":")
        uid = int(parts[1])
        main_cat = parts[2]
        sub_cat = parts[3]
        qty = int(parts[4])
        total_price = float(parts[5])
        
        # Check stock first
        stock_count = count_items(main_cat, sub_cat)

        if stock_count < qty:
            # Not enough stock, show override option
            keyboard = InlineKeyboardMarkup([
                [InlineKeyboardButton("✅ Force Confirm", callback_data=f"force_confirm:{uid}:{main_cat}:{sub_cat}:{qty}:{total_price}"),
                 InlineKeyboardButton("❌ Cancel", callback_data=f"cancel_manual:{uid}")]
            ])
            await query.edit_message_caption(query.message.caption + "\n\n⚠️ স্টকে পর্যাপ্ত আইটেম নেই। আপনি কি নিশ্চিত করতে চান?\n\nযদি নিশ্চিত করেন, আপনাকে ম্যানুয়ালি আইটেমটি পাঠাতে হবে।", reply_markup=keyboard)
            return
            
        # Manually pop items from excel
        items = pop_items_from_excel(main_cat, sub_cat, qty)
        
        total_sales += total_price
        sales_count_per_category[sub_cat] = sales_count_per_category.get(sub_cat, 0) + qty
        transaction_log.append(('sale', uid, total_price, time.time()))
        user_sales[uid] = user_sales.get(uid, 0) + total_price
        save_user_data()

        await context.bot.send_message(
            chat_id=uid, 
            text="✅ আপনার অর্ডার নিশ্চিত করা হয়েছে। অ্যাডমিন শীঘ্রই আপনাকে বিস্তারিত তথ্য পাঠাবেন।"
        )
        await query.edit_message_caption(query.message.caption + f"\n\n✅ Confirmed by Admin. User will be contacted manually.", reply_markup=None)

    elif data.startswith("force_confirm:"):
        parts = data.split(":")
        uid = int(parts[1])
        main_cat = parts[2]
        sub_cat = parts[3]
        qty = int(parts[4])
        total_price = float(parts[5])

        # Force confirm even without stock
        total_sales += total_price
        sales_count_per_category[sub_cat] = sales_count_per_category.get(sub_cat, 0) + qty
        transaction_log.append(('sale', uid, total_price, time.time()))
        user_sales[uid] = user_sales.get(uid, 0) + total_price
        save_user_data()
        
        # Admin must manually deliver, so no items are popped.
        await context.bot.send_message(
            chat_id=uid, 
            text="✅ আপনার অর্ডার নিশ্চিত করা হয়েছে। অ্যাডমিন শীঘ্রই আপনাকে বিস্তারিত তথ্য পাঠাবেন।"
        )
        await query.edit_message_caption(query.message.caption + f"\n\n✅ Force Confirmed by Admin. User will be contacted manually.", reply_markup=None)

    elif data.startswith("cancel_manual:"):
        _, uid = data.split(":")
        
        await context.bot.send_message(
            chat_id=uid,
            text="❌ দুঃখিত, আপনার অর্ডারটি বাতিল করা হয়েছে।"
        )
        await query.edit_message_caption(query.message.caption + "\n\n❌ Cancelled by Admin", reply_markup=None)
        
    else: # Default cancel
        _, uid = data.split(":")
        
        await context.bot.send_message(
            chat_id=uid,
            text="❌ দুঃখিত, আপনার অর্ডারটি বাতিল করা হয়েছে।"
        )
        
        await query.edit_message_caption(query.message.caption + "\n\n❌ Cancelled by Admin", reply_markup=None)


async def admin_deposit_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    global total_deposits, transaction_log, user_deposits, balances
    query = update.callback_query
    await query.answer() # বাটন ক্লিকের লোডিং স্টেট সরিয়ে দেয়
    data = query.data

    # শুধুমাত্র অ্যাডমিনই এই অ্যাকশনটি করতে পারবে
    if update.effective_user.id != ADMIN_ID:
        try:
            # Unauthorized হলে বাটন রিমুভ করে দেওয়া উচিত
            await query.edit_message_caption("❌ Unauthorized.", reply_markup=None)
        except:
            pass
        return
    
    # === ডিপোজিট নিশ্চিতকরণ (Confirm) ব্লক ===
    if data.startswith("deposit_confirm:"):
        parts = data.split(":")
        uid = int(parts[1])
        amount = float(parts[2])
        
        # ব্যালেন্স আপডেট করা হলো
        balances[uid] = balances.get(uid, 0) + amount
        
        # অন্যান্য ডেটা আপডেট ও সেভ করা হলো
        total_deposits += amount
        transaction_log.append(('deposit', uid, amount, time.time()))
        user_deposits[uid] = user_deposits.get(uid, 0) + amount
        save_user_data()

        # ব্যবহারকারীকে সফলতার মেসেজ
        await context.bot.send_message(
            chat_id=uid,
            text=f"✅ আপনার ডিপোজিট সফল হয়েছে। আপনার ব্যালেন্সে {amount} টাকা যোগ হয়েছে।\n"
                 f"নতুন ব্যালেন্স: {balances[uid]} টাকা।"
        )
        
        # অ্যাডমিনের মেসেজ আপডেট করা হলো (নতুন মোট ব্যালেন্স সহ)
        try:
            await query.edit_message_caption(
                # ক্যাপশনের শেষে নতুন টেক্সট যোগ করা হচ্ছে
                query.message.caption + 
                f"\n\n✅ নিশ্চিত করা হয়েছে। {amount} টাকা ইউজার <code>{uid}</code> এর ব্যালেন্সে যোগ করা হয়েছে। "
                f"<b>বর্তমান মোট ব্যালেন্স: {balances[uid]} টাকা।</b>", 
                reply_markup=None,
                parse_mode='HTML' # ত্রুটি এড়াতে এবং বোল্ড করার জন্য HTML ব্যবহার করা হয়েছে
            )
        except Exception as e:
            logger.error(f"Failed to edit message caption on deposit confirmation: {e}")
            # ক্যাপশন আপডেট না হলে অ্যাডমিনকে বিকল্প মেসেজ পাঠানো হবে
            await context.bot.send_message(
                chat_id=ADMIN_ID, 
                text=f"⚠️ Failed to update deposit confirmation message. User ID: {uid}, Amount: {amount}"
            )
            
    # === ডিপোজিট বাতিলকরণ (Cancel) ব্লক ===
    else: # deposit_cancel:
        _, uid = data.split(":")
        uid = int(uid)
        
        # ব্যবহারকারীকে বাতিলের মেসেজ
        await context.bot.send_message(
            chat_id=uid,
            text="❌ দুঃখিত, আপনার ডিপোজিট রিকোয়েস্ট বাতিল করা হয়েছে।"
        )
        
        # অ্যাডমিনের মেসেজ আপডেট করা হলো
        try:
            await query.edit_message_caption(
                query.message.caption + "\n\n❌ Deposit cancelled by Admin", 
                reply_markup=None
            )
        except Exception as e:
            logger.error(f"Failed to edit message caption on deposit cancellation: {e}")
            await context.bot.send_message(
                chat_id=ADMIN_ID, 
                text=f"⚠️ Failed to update deposit cancellation message. User ID: {uid}"
            )
# ================ MAIN =================
def main():
    app = ApplicationBuilder().token(BOT_TOKEN).build()

    conv = ConversationHandler(
        entry_points=[
            CommandHandler("start", start),
            MessageHandler(filters.TEXT & ~filters.COMMAND, menu_handler)
        ],
        states={
            BUY_MENU: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & ~filters.COMMAND, user_choose_category)
            ],
            BUY_SUB_MENU: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Categories"), back_to_categories_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & ~filters.COMMAND, user_choose_subcategory)
            ],
            GET_QUANTITY: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Sub-Categories"), back_to_subcategories_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_quantity)
            ],
            WAIT_SCREENSHOT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Sub-Categories"), back_to_subcategories_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & filters.Regex("✅ Confirm Purchase"), user_send_screenshot),
                MessageHandler(filters.PHOTO | filters.TEXT, user_send_screenshot),
            ],
            ADMIN_PANEL: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & ~filters.COMMAND, admin_panel_handler)
            ],
            MANAGE_CATEGORY: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, manage_category_handler)
            ],
            MANAGE_SUB_CATEGORY: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Manage Main Categories"), back_to_manage_main_categories_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, manage_sub_category_handler)
            ],
            ADD_MAIN_CAT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, add_main_category)
            ],
            REMOVE_MAIN_CAT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, remove_main_category)
            ],
            ADD_SUB_CAT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Manage Main Categories"), back_to_manage_main_categories_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, add_sub_category)
            ],
            REMOVE_SUB_CAT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Manage Main Categories"), back_to_manage_main_categories_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, remove_sub_category)
            ],
            ADD_ITEMS: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & filters.Regex("🔙 Manage Main Categories"), back_to_manage_main_categories_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, add_item_line)
            ],
            EDIT_PAYMENT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, edit_payment_info)
            ],
            
            EDIT_PRICE_MAIN: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, edit_price_main_handler)
            ],
            EDIT_PRICE_SUB: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, edit_price_sub_handler)
            ],
            RECEIVE_NEW_PRICE: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_price)
            ],
            DEPOSIT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.TEXT & ~filters.COMMAND, deposit_handler),
            ],
            GET_DEPOSIT_AMOUNT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 Back to Main Menu"), start),
                MessageHandler(filters.PHOTO, receive_deposit_screenshot)
            ],
            DASHBOARD: [
                MessageHandler(filters.TEXT & filters.Regex("🔄 ড্যাশবোর্ড রিফ্রেশ করুন"), handle_dashboard_refresh),
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler)
            ],
            SEND_NOTICE: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, send_notice_text)
            ],
            SEARCH_USER_PROFILE: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, search_and_show_user_profile)
            ],
            MANAGE_PAYMENT_CATEGORIES: [
                CallbackQueryHandler(toggle_payment_method),
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler)
            ],
            # --- NEW STATES HANDLERS ---
            SEARCH_USER_FOR_BALANCE: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, search_user_for_balance)
            ],
            BALANCE_EDIT_ACTION: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, balance_edit_action_handler)
            ],
            RECEIVE_BALANCE_EDIT_AMOUNT: [
                MessageHandler(filters.TEXT & filters.Regex("🔙 অ্যাডমিন প্যানেল"), back_to_admin_panel_handler),
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_balance_edit_amount)
            ]
        },
        fallbacks=[CommandHandler("start", start)]
    )

    
    app.add_handler(conv)
    app.add_handler(CallbackQueryHandler(admin_order_action, pattern="^(confirm:|cancel:|confirm_manual:|cancel_manual:|force_confirm:)"))
    app.add_handler(CallbackQueryHandler(admin_deposit_action, pattern="^(deposit_confirm:|deposit_cancel:)"))
    app.add_handler(CallbackQueryHandler(toggle_payment_method, pattern="^toggle_payment:"))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND & filters.Regex("🔙 Back to Main Menu"), start))
    
    logger.info("🤖 Bot running...")
    app.run_polling()

if __name__ == "__main__":
    main()
